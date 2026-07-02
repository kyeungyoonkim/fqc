from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from src.config import ExcelConfig
from src.excel_writer import append_daily_row


def _make_workbook(tmp_path: Path, with_table: bool) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    headers = ["Date", "CTV E03", "CTV E19", "CTV T11", "CTV 전체"]
    ws.append(headers)
    ws.append([date(2026, 6, 30), 1.1, 2.2, 3.3, 4.4])

    if with_table:
        table = Table(displayName="FQC_Data", ref="A1:E2")
        ws.add_table(table)

    path = tmp_path / "tracking.xlsx"
    wb.save(path)
    return path


def _excel_config(path: Path, table_name: str | None) -> ExcelConfig:
    return ExcelConfig(
        path=str(path),
        sheet_name="Data",
        table_name=table_name,
        date_column_header="Date",
        column_headers={
            "CTV_E03": "CTV E03",
            "CTV_E19": "CTV E19",
            "CTV_T11": "CTV T11",
            "CTV_OVERALL": "CTV 전체",
        },
    )


def test_append_daily_row_appends_after_last_used_row(tmp_path):
    path = _make_workbook(tmp_path, with_table=False)
    cfg = _excel_config(path, table_name=None)

    append_daily_row(
        cfg,
        date(2026, 7, 1),
        {"CTV_E03": 0.9, "CTV_E19": 1.5, "CTV_T11": 2.0, "CTV_OVERALL": 1.3},
    )

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Data"]

    assert ws.cell(row=3, column=1).value.date() == date(2026, 7, 1)
    assert ws.cell(row=3, column=2).value == 0.9
    assert ws.cell(row=3, column=3).value == 1.5
    assert ws.cell(row=3, column=4).value == 2.0
    assert ws.cell(row=3, column=5).value == 1.3

    # existing row untouched
    assert ws.cell(row=2, column=1).value.date() == date(2026, 6, 30)


def test_append_daily_row_extends_table_ref(tmp_path):
    path = _make_workbook(tmp_path, with_table=True)
    cfg = _excel_config(path, table_name="FQC_Data")

    append_daily_row(
        cfg,
        date(2026, 7, 1),
        {"CTV_E03": 0.9, "CTV_E19": 1.5, "CTV_T11": 2.0, "CTV_OVERALL": 1.3},
    )

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Data"]
    table = ws.tables["FQC_Data"]
    assert table.ref == "A1:E3"


def test_append_daily_row_raises_on_missing_header(tmp_path):
    path = _make_workbook(tmp_path, with_table=False)
    cfg = _excel_config(path, table_name=None)
    cfg.column_headers["CTV_OVERALL"] = "Nonexistent Header"

    try:
        append_daily_row(
            cfg,
            date(2026, 7, 1),
            {"CTV_E03": 0.9, "CTV_OVERALL": 1.3},
        )
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "Nonexistent Header" in str(exc)


def test_append_daily_row_fills_multiple_consecutive_days(tmp_path):
    path = _make_workbook(tmp_path, with_table=False)
    cfg = _excel_config(path, table_name=None)

    for i, d in enumerate([date(2026, 7, 1), date(2026, 7, 2), date(2026, 7, 3)]):
        append_daily_row(
            cfg,
            d,
            {"CTV_E03": i * 0.1, "CTV_E19": i * 0.2, "CTV_T11": i * 0.3, "CTV_OVERALL": i * 0.4},
        )

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Data"]
    assert ws.cell(row=5, column=1).value.date() == date(2026, 7, 3)
    assert ws.cell(row=5, column=2).value == 0.2
